from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT_DIR = Path("completed_docs")
SONG = "宋体"
HEI = "黑体"


def set_run_font(run, name=SONG, size=12, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(18)
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def setup_doc(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = SONG
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.first_line_indent = Cm(0.74)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(20)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    for name, size, font, center in [
        ("Title", 22, SONG, True),
        ("Heading 1", 14, HEI, True),
        ("Heading 2", 12, HEI, False),
        ("Heading 3", 12, HEI, False),
    ]:
        style = styles[name]
        style.font.name = font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.font.size = Pt(size)
        style.font.bold = True if font == HEI else False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(20)
        style.paragraph_format.space_before = Pt(6 if name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(6 if name.startswith("Heading") else 0)
        style.paragraph_format.first_line_indent = None if center else Cm(0.74)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT


def add_para(doc, text="", style=None, align=None, first=True, size=12, bold=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(20)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0.74) if first else None
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    return p


def add_title(doc, title):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    r = p.add_run(title)
    set_run_font(r, size=22, bold=True)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run_font(r, HEI, 14 if level == 1 else 12, bold=True)
    return p


def add_table(doc, caption, headers, rows, widths=None):
    add_para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, first=False, size=10.5)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(hdr[i], "D9EAF7")
        if widths:
            hdr[i].width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], str(val), align=WD_ALIGN_PARAGRAPH.CENTER if len(str(val)) < 12 else WD_ALIGN_PARAGRAPH.LEFT)
            if widths:
                cells[i].width = Cm(widths[i])
    add_para(doc, "", first=False)
    return table


def add_ascii_diagram(doc, caption, lines):
    add_para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, first=False, size=10.5)
    for line in lines:
        p = add_para(doc, line, first=False, size=10.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_experiment_header(doc, title):
    add_title(doc, title)
    add_para(doc, "姓名：          班级：          学号：             成绩：", first=False)


def experiment2():
    doc = Document()
    setup_doc(doc)
    add_experiment_header(doc, "实验2 智能系统系统设计实验")
    add_heading(doc, "一、实验要求", 1)
    add_para(doc, "完成人工智能项目的系统设计、项目功能设计、数据库设计和系统技术实现架构设计。本实验以“智学伴 AI 个性化学习伴侣系统”为对象，围绕用户认证、知识库管理、文档处理、向量检索、RAG 问答和学习记录等功能进行设计。")
    add_heading(doc, "二、实验目的", 1)
    for item in ["学习项目系统功能设计方法，能够从需求出发拆分功能模块。", "学会根据业务对象进行数据库实体、属性和关系设计。", "掌握 Web 智能系统的前后端架构、接口划分和 RAG 技术实现流程。"]:
        add_para(doc, item)
    add_heading(doc, "三、实验内容", 1)
    add_heading(doc, "1. 系统总体功能结构设计", 2)
    add_para(doc, "系统面向学生用户提供基于资料库的智能问答服务。用户登录后可创建知识库、上传 PDF/TXT/DOCX 文档，系统完成文本抽取、分块、向量化和 Faiss 索引构建；提问时先检索知识库，再将相关片段注入提示词并调用 Ollama 或 API 模型生成回答，同时保存学习记录。")
    add_ascii_diagram(doc, "图2-1 系统总体功能结构图", [
        "智学伴智能学习系统",
        "├─ 用户认证：注册、登录、邮箱验证、密码重置、账户锁定",
        "├─ 个人中心：资料维护、头像上传、模型/API Key 配置",
        "├─ 知识库管理：创建、重命名、删除、统计、文件来源管理",
        "├─ 文档处理：PDF 文本抽取、OCR、DOCX/TXT 读取、文本分块",
        "├─ 智能问答：向量检索、重排、提示词构造、LLM 生成",
        "└─ 学习记录：会话分组、历史查询、反馈标记"
    ])
    add_table(doc, "表2-1 系统功能模块设计", ["模块", "主要功能", "关键文件"], [
        ["用户认证模块", "注册、登录、JWT、CSRF、防暴力破解、日志记录", "app/routes/auth.py、app/services/auth_service.py"],
        ["知识库模块", "知识库实例管理、文档入库、向量列表、来源删除", "app/services/knowledge_service.py"],
        ["文档处理模块", "PDF/OCR、DOCX、TXT 文本抽取", "app/utils/file_processor.py"],
        ["向量检索模块", "嵌入生成、向量归一化、Faiss 索引、持久化", "app/utils/vector_service.py"],
        ["AI 问答模块", "Ollama/DeepSeek 调用、RAG 提示词生成、回答返回", "app/services/ai_service.py"],
        ["前端交互模块", "登录注册、首页问答、知识库页面、个人中心", "templates、static/js"]
    ], [2.8, 8.0, 5.0])
    add_heading(doc, "2. 数据库实体关系与表设计", 2)
    add_para(doc, "数据库采用 MySQL，ORM 使用 Flask-SQLAlchemy。核心实体包括用户、学习记录和用户操作日志。用户与学习记录是一对多关系，用户操作日志记录登录、注册、邮箱验证和密码重置等行为。知识库向量数据采用本地 pickle 文件与 Faiss 索引文件持久化，文件名以 knowledge_base_{name} 区分。")
    add_ascii_diagram(doc, "图2-2 数据库 ER 图", [
        "User(用户) 1 ───── n LearningRecord(学习记录)",
        "User(用户) 1 ───── n UserLog(用户操作日志)",
        "KnowledgeBase(知识库文件/索引) ── VectorStore(Faiss + 元数据)"
    ])
    add_table(doc, "表2-2 users 用户表设计", ["字段", "类型", "说明"], [
        ["id", "Integer PK", "用户主键"],
        ["username", "String(50)", "用户名，唯一且非空"],
        ["password", "String(255)", "bcrypt 加密后的密码"],
        ["email", "String(100)", "邮箱，支持验证和密码重置"],
        ["login_attempts", "Integer", "登录失败次数"],
        ["locked_until", "DateTime", "账户锁定截止时间"],
        ["deepseek_api_key/kimi_api_key/zhipu_api_key", "String(255)", "用户自定义模型密钥"],
        ["use_ollama/ollama_model", "Boolean/String", "本地模型开关和模型名称"]
    ], [4.0, 4.0, 8.0])
    add_table(doc, "表2-3 learning_records 学习记录表设计", ["字段", "类型", "说明"], [
        ["id", "Integer PK", "记录主键"],
        ["user_id", "Integer FK", "关联 users.id"],
        ["question", "Text", "用户问题"],
        ["ai_answer", "Text", "AI 回答"],
        ["sources", "Text", "检索来源 JSON"],
        ["session_id", "String(50)", "会话分组编号"],
        ["knowledge_base", "String(100)", "使用的知识库名称"],
        ["created_at", "DateTime", "创建时间，建立索引"]
    ], [4.0, 4.0, 8.0])
    add_heading(doc, "3. 技术实现架构设计", 2)
    add_para(doc, "系统采用 Browser/Server 架构。前端由 Jinja2 模板、HTML、CSS 与 JavaScript 组成，负责页面展示、表单提交和接口调用；后端使用 Flask 提供页面路由与 REST API；业务层封装认证、邮件、AI 问答和知识库服务；数据层由 MySQL、Faiss 本地索引和上传文件目录组成。")
    add_ascii_diagram(doc, "图2-3 前后端与 RAG 技术架构图", [
        "浏览器页面 → Flask 路由/API → Service 业务层 → 数据/模型层",
        "上传文档 → FileProcessor → TextSplitter → EmbeddingModel → VectorStore(Faiss)",
        "用户提问 → 向量检索 → Reranker → Prompt → Ollama/DeepSeek → 回答与来源"
    ])
    add_table(doc, "表2-4 主要接口设计", ["接口", "方法", "功能"], [
        ["/register、/login", "GET/POST", "注册与登录页面及认证逻辑"],
        ["/api/records/sessions", "GET", "查询用户最近会话列表"],
        ["/api/records/session/<id>", "GET", "查询单个会话详情"],
        ["/api/feedback", "POST", "提交回答是否有帮助的反馈"],
        ["/knowledge-base", "GET", "知识库管理页面"],
        ["知识库相关 API", "POST/GET/DELETE", "文档上传、知识库统计、向量来源查看和删除"]
    ], [5.2, 2.5, 8.3])
    add_heading(doc, "四、结论", 1)
    add_para(doc, "通过本次系统设计实验，完成了智学伴系统的功能分解、数据库建模和技术架构设计。系统设计体现了 Web 应用与人工智能能力的结合：用户侧关注学习资料管理和问答体验，后端侧通过文档处理、向量检索、重排和大模型生成实现 RAG 流程，为后续开发和测试奠定了结构清晰的基础。")
    return doc


def experiment3():
    doc = Document()
    setup_doc(doc)
    add_experiment_header(doc, "实验3 智能系统开发实验")
    add_heading(doc, "一、实验要求", 1)
    add_para(doc, "完成人工智能项目的开发实现，说明各个功能模块的实现思路、关键流程和核心代码组织。")
    add_heading(doc, "二、实验目的", 1)
    for item in ["学习智能系统功能模块的代码实现方法。", "掌握 Flask Web 项目、数据库模型、服务层和前端交互的开发流程。", "理解 RAG 智能问答链路在实际项目中的实现方式。"]:
        add_para(doc, item)
    add_heading(doc, "三、实验内容", 1)
    add_heading(doc, "1. 项目工程结构实现", 2)
    add_para(doc, "项目以 app.py 和 run.py 作为入口，config.py 管理环境变量与数据库连接，app 目录按 models、routes、services、utils 进行分层，templates 与 static 分别存放前端页面和静态资源。该结构使页面路由、业务逻辑、数据模型和工具函数保持相对独立，便于维护。")
    add_table(doc, "表3-1 工程目录与职责", ["目录/文件", "职责"], [
        ["app.py", "创建 Flask 应用、注册路由、定义核心 API 与页面入口"],
        ["run.py", "初始化数据库并启动服务"],
        ["app/models", "定义 User、LearningRecord、UserLog 数据模型"],
        ["app/routes", "处理认证相关 Blueprint 路由"],
        ["app/services", "封装认证、邮件、AI、知识库、重排等业务服务"],
        ["app/utils", "封装文档处理、文本分块、向量存储等工具"],
        ["templates/static", "实现登录注册、首页问答、知识库和个人中心页面"]
    ], [5.0, 11.0])
    add_heading(doc, "2. 用户认证模块实现", 2)
    add_para(doc, "认证模块通过 bcrypt 对密码加密，使用 JWT 保存登录状态，并结合 CSRF Token 防止跨站请求伪造。系统设置密码复杂度校验规则，连续登录失败达到 5 次后锁定账号 15 分钟，同时写入 UserLog 便于审计。邮箱验证码和密码重置令牌通过随机安全字符串生成。")
    add_ascii_diagram(doc, "图3-1 用户登录实现流程图", [
        "输入用户名/邮箱和密码",
        "↓",
        "查询用户并检查 locked_until",
        "↓",
        "bcrypt 校验密码",
        "↓ 成功                         ↓ 失败",
        "生成 JWT 与 CSRF Token          login_attempts + 1",
        "写入登录日志并进入首页           达到阈值则锁定账户"
    ])
    add_heading(doc, "3. 文档入库与知识库实现", 2)
    add_para(doc, "知识库模块以 QAModule 为核心，每个知识库名称对应一个单例实例。用户上传文档后，FileProcessor 根据扩展名选择 PDF、DOCX 或 TXT 处理方式；PDF 优先使用 PyMuPDF 抽取文本，当页面文本不足时自动调用 PaddleOCR；文本经 TextSplitter 按页面和段落切分，默认块大小 500 字、重叠 100 字。")
    add_ascii_diagram(doc, "图3-2 文档入库实现流程图", [
        "上传文件",
        "↓",
        "FileProcessor 抽取文本",
        "↓",
        "TextSplitter 分块",
        "↓",
        "EmbeddingModel 批量生成向量",
        "↓",
        "VectorStore 写入 Faiss 索引与元数据文件"
    ])
    add_heading(doc, "4. RAG 问答模块实现", 2)
    add_para(doc, "用户提问时，系统先将问题转为查询向量，在 Faiss 中召回相似片段；若启用重排服务，则进一步按相关性保留前 5 个片段。随后 _build_prompt 将上下文和问题组合为提示词，调用 Ollama 的 /api/chat 或 API 模型生成回答，并返回答案、来源片段和 context_used 标记。")
    add_table(doc, "表3-2 RAG 关键类与方法", ["类/方法", "实现作用"], [
        ["QAModule.query_with_knowledge", "完成查询向量生成、召回、重排、提示词构造和答案返回"],
        ["EmbeddingModel.get_embedding", "调用 bge-m3 嵌入模型；连接失败时使用字符哈希向量降级"],
        ["VectorStore.search", "归一化查询向量并使用 Faiss IndexFlatIP 检索"],
        ["RerankerService.rerank_with_metadata", "对召回结果重新排序并保留元数据"],
        ["AIService.ask_question", "在 Ollama、本地配置和远程 API 之间选择模型调用方式"]
    ], [5.5, 10.5])
    add_heading(doc, "5. 前端交互实现", 2)
    add_para(doc, "前端通过模板页面展示登录注册、首页问答、知识库管理和个人中心。JavaScript 负责提交表单、维护会话状态、调用问答和知识库接口，并把回答、来源和历史记录渲染到页面中。页面与后端接口分离，使系统可以逐步扩展为更完整的前后端分离架构。")
    add_heading(doc, "四、结论", 1)
    add_para(doc, "通过本次开发实验，项目完成了从用户认证到智能问答的主要功能闭环。系统代码按模型层、服务层、工具层和页面层组织，RAG 功能具备文档处理、向量化、检索、重排、模型生成和记录保存等完整步骤，能够支撑后续功能测试、性能优化和界面完善。")
    return doc


def experiment4():
    doc = Document()
    setup_doc(doc)
    add_experiment_header(doc, "实验4 智能系统系统测试实验")
    add_heading(doc, "一、实验要求", 1)
    add_para(doc, "完成人工智能项目的测试，对各功能模块进行功能测试和性能测试，验证系统是否满足设计目标。")
    add_heading(doc, "二、实验目的", 1)
    for item in ["学习智能系统功能测试方法，能够围绕业务流程设计测试用例。", "学会对系统性能、异常处理和安全控制进行测试。", "掌握 RAG 问答系统中文档入库、检索和回答质量的验证方法。"]:
        add_para(doc, item)
    add_heading(doc, "三、实验内容", 1)
    add_heading(doc, "1. 测试环境", 2)
    add_table(doc, "表4-1 测试环境配置", ["项目", "配置"], [
        ["操作系统", "Windows，本地开发环境"],
        ["后端框架", "Python + Flask 3.0.0 + Flask-SQLAlchemy"],
        ["数据库", "MySQL，字符集 utf8mb4"],
        ["向量检索", "Faiss CPU，IndexFlatIP"],
        ["文档处理", "PyMuPDF、python-docx、PaddleOCR"],
        ["模型服务", "Ollama qwen2.5:7b、bge-m3；可选 DeepSeek/Kimi/智谱 API"],
        ["浏览器", "Chrome/Edge 等现代浏览器"]
    ], [4.0, 12.0])
    add_heading(doc, "2. 功能测试", 2)
    add_table(doc, "表4-2 功能测试用例", ["编号", "测试项", "输入/操作", "预期结果"], [
        ["F01", "用户注册", "输入合法用户名、邮箱和复杂密码", "发送验证码并创建待验证用户"],
        ["F02", "登录认证", "输入正确账号密码", "生成 Token，跳转首页"],
        ["F03", "账户锁定", "连续输入错误密码 5 次", "账户锁定 15 分钟并记录日志"],
        ["F04", "知识库创建", "创建名为 default/course 的知识库", "知识库实例和上传目录可用"],
        ["F05", "文档上传", "上传 PDF、TXT、DOCX 文件", "文本被抽取并写入向量库"],
        ["F06", "智能问答", "基于上传资料提问", "返回答案、来源片段和相似度信息"],
        ["F07", "历史记录", "完成多轮问答后查看记录", "按 session_id 分组显示最近会话"],
        ["F08", "反馈提交", "点击有用/无用", "更新 learning_records.is_helpful"]
    ], [1.5, 3.0, 5.5, 6.0])
    add_heading(doc, "3. RAG 链路测试", 2)
    add_para(doc, "RAG 链路测试重点验证“资料是否能进入知识库”“提问是否能召回相关片段”“回答是否基于上下文”。测试时可准备一份包含课程知识点的 DOCX 文档，上传后提问文档中出现的概念；若返回 context_used=true 且 sources 列表包含对应片段，说明检索增强链路有效。")
    add_ascii_diagram(doc, "图4-1 RAG 测试流程图", [
        "准备测试文档 → 上传入库 → 查看向量数量",
        "↓",
        "提出命中文档内容的问题 → 检查 sources 与 answer",
        "↓",
        "提出无关问题 → 检查系统是否说明上下文不足"
    ])
    add_heading(doc, "4. 性能测试", 2)
    add_table(doc, "表4-3 性能测试指标", ["测试项", "方法", "期望结果"], [
        ["启动性能", "运行 python run.py 并访问首页", "服务正常启动，页面可访问"],
        ["文档入库耗时", "上传 1MB、5MB 文档并记录耗时", "小文件数秒内完成，大文件可接受等待并无异常"],
        ["向量检索耗时", "知识库中存在多段文本时连续提问", "Faiss 检索保持较低延迟"],
        ["并发访问", "多个浏览器会话同时登录和提问", "会话互不串扰，接口无 500 错误"],
        ["缓存效果", "重复访问会话列表", "10 分钟缓存降低数据库查询压力"]
    ], [3.5, 6.0, 6.5])
    add_heading(doc, "5. 异常与安全测试", 2)
    add_table(doc, "表4-4 异常与安全测试", ["测试项", "测试方法", "预期结果"], [
        ["不支持文件格式", "上传 exe 或未知扩展名文件", "返回不支持格式错误"],
        ["大文件限制", "上传超过 MAX_CONTENT_LENGTH 的文件", "返回 413 文件过大提示"],
        ["模型不可用", "关闭 Ollama 后提问", "给出无法连接或配置提示，不导致服务崩溃"],
        ["CSRF 防护", "缺少 X-CSRF-Token 提交修改请求", "返回 403"],
        ["Token 过期", "使用过期 JWT 访问受保护页面", "跳转登录或返回 401"]
    ], [3.5, 6.0, 6.5])
    add_heading(doc, "6. 测试结论", 2)
    add_para(doc, "从测试设计结果看，系统主要功能链路完整，认证、知识库、文档处理、向量检索、问答生成和记录保存均有明确测试点。后续优化方向包括：增加自动化接口测试覆盖率，使用更真实的数据集评估回答准确率，并对大文件 OCR、向量重建和并发问答场景进行更细粒度的性能压测。")
    add_heading(doc, "四、结论", 1)
    add_para(doc, "通过本次系统测试实验，完成了智学伴项目的功能测试、RAG 链路测试、性能测试和异常安全测试设计。测试结果能够验证系统是否达到课程设计目标，也为后续修复缺陷、优化性能和完善用户体验提供依据。")
    return doc


def course_report():
    doc = Document()
    setup_doc(doc)
    add_title(doc, "智学伴智能学习系统")
    add_para(doc, "", first=False)
    for line in [
        "课程名称：智能系统应用实践",
        "设计题目：智学伴智能学习系统",
        "学院：计算机学院",
        "专业班级：____________________",
        "学生姓名：____________________",
        "学号：____________________",
        "指导教师：____________________",
        "完成日期：2026年6月",
    ]:
        p = add_para(doc, line, first=False, size=14)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    add_heading(doc, "内容摘要", 1)
    add_para(doc, "本课程设计完成了一个基于 Flask 的智能学习伴侣系统。系统支持用户注册登录、邮箱验证、个人模型配置、文档上传、知识库管理、向量检索和 RAG 智能问答。后端利用 PyMuPDF、PaddleOCR 和 python-docx 完成多格式文档解析，使用 bge-m3 生成文本向量并通过 Faiss 建立索引，结合 Ollama 或远程模型生成回答。设计过程加深了对 Web 智能系统分层架构、数据库设计和检索增强生成技术的理解。")
    add_para(doc, "关键词：智能学习；Flask；知识库；向量检索；RAG")
    doc.add_page_break()

    add_heading(doc, "目录", 1)
    contents = [
        "第一章 绪论",
        "第二章 需求分析",
        "第三章 系统设计",
        "第四章 系统实现",
        "第五章 系统测试",
        "第六章 总结",
        "参考文献",
        "评定意见页",
    ]
    for item in contents:
        add_para(doc, item + " ........................................", first=False)
    doc.add_page_break()

    add_heading(doc, "第一章 绪论", 1)
    add_heading(doc, "1.1 设计背景", 2)
    add_para(doc, "随着生成式人工智能和大语言模型的发展，学习辅助系统不再局限于固定题库和静态资料展示。学生在课程学习中常常需要对课件、实验文档和参考资料进行检索、总结和问答，如果仅依赖普通聊天模型，容易出现回答与课程资料不一致的问题。因此，本设计选择“智学伴智能学习系统”作为课程设计题目，通过知识库检索增强生成方式，让系统优先基于用户上传资料回答问题。")
    add_heading(doc, "1.2 设计目的与意义", 2)
    add_para(doc, "本系统的设计目的在于实现一个可运行的 Web 智能系统，综合运用 Flask 后端开发、MySQL 数据库、文档解析、向量检索和大模型调用等知识。系统能够帮助学生沉淀个人学习资料，通过自然语言提问快速获得资料相关答案，同时保存学习记录，便于复习和追踪。")
    add_heading(doc, "1.3 开发环境", 2)
    add_table(doc, "表1-1 开发环境", ["类别", "内容"], [
        ["开发语言", "Python、HTML、CSS、JavaScript"],
        ["后端框架", "Flask、Flask-CORS、Flask-SQLAlchemy"],
        ["数据库", "MySQL、PyMySQL"],
        ["智能组件", "Ollama、qwen2.5:7b、bge-m3、DeepSeek API 可选"],
        ["文档处理", "PyMuPDF、PaddleOCR、python-docx"],
        ["向量检索", "numpy、faiss-cpu"],
        ["运行方式", "配置 .env 后执行 python run.py，默认访问 http://localhost:5000"]
    ], [4.0, 12.0])

    add_heading(doc, "第二章 需求分析", 1)
    add_heading(doc, "2.1 功能需求", 2)
    add_para(doc, "系统主要面向学生个人学习场景，核心需求包括：用户可安全注册和登录；可维护个人资料和模型配置；可创建知识库并上传课程文档；系统可自动解析文档并建立向量索引；用户可围绕资料内容提问并获得带来源的回答；系统可保存历史会话和用户反馈。")
    add_table(doc, "表2-1 功能需求表", ["需求编号", "需求名称", "说明"], [
        ["R1", "用户认证", "注册、登录、邮箱验证、密码重置、账户锁定"],
        ["R2", "个人中心", "头像上传、专业年级维护、API Key 和本地模型配置"],
        ["R3", "知识库管理", "创建、删除、重命名、统计和文件来源管理"],
        ["R4", "文档处理", "支持 PDF、TXT、DOCX，PDF 文本不足时使用 OCR"],
        ["R5", "智能问答", "基于向量检索和大模型生成答案，返回来源片段"],
        ["R6", "学习记录", "保存问题、回答、来源、会话编号和反馈结果"]
    ], [2.0, 4.0, 10.0])
    add_heading(doc, "2.2 非功能需求", 2)
    add_para(doc, "系统需要具备可维护性、安全性和一定的性能。可维护性体现在模块分层清晰；安全性体现在密码加密、JWT、CSRF 和登录失败锁定；性能体现在使用 Faiss 进行本地向量检索，并对会话列表结果设置 10 分钟缓存。系统还应具备容错能力，当 Ollama 不可用时给出明确提示，嵌入模型失败时使用本地字符哈希向量降级。")

    add_heading(doc, "第三章 系统设计", 1)
    add_heading(doc, "3.1 总体架构设计", 2)
    add_para(doc, "系统采用分层 B/S 架构。前端模板负责页面展示和用户交互，Flask 路由负责接收请求，服务层负责认证、AI 问答和知识库业务，工具层负责文档解析、文本分块和向量检索，数据层包括 MySQL、上传目录、Faiss 索引和向量元数据文件。")
    add_ascii_diagram(doc, "图3-1 系统总体架构图", [
        "浏览器层：登录注册、首页问答、知识库、个人中心",
        "应用层：Flask 路由、REST API、JWT/CSRF",
        "业务层：AuthService、AIService、QAModule、EmailService",
        "算法层：FileProcessor、TextSplitter、EmbeddingModel、VectorStore、Reranker",
        "数据层：MySQL、uploads、Faiss index、vector_store.pkl"
    ])
    add_heading(doc, "3.2 数据库设计", 2)
    add_para(doc, "数据库设计围绕用户和学习记录展开。users 表保存认证信息、邮箱状态、头像、专业年级和模型配置；learning_records 表保存问答历史、来源、会话编号和反馈；user_logs 表保存用户操作审计信息。")
    add_table(doc, "表3-1 主要数据表", ["表名", "作用", "关键字段"], [
        ["users", "保存用户账号与个性化配置", "id、username、password、email、use_ollama、ollama_model"],
        ["learning_records", "保存问答历史与反馈", "user_id、question、ai_answer、sources、session_id、knowledge_base"],
        ["user_logs", "保存注册、登录、密码重置等日志", "user_id、action、ip_address、details、status、created_at"]
    ], [3.5, 5.5, 7.0])
    add_heading(doc, "3.3 RAG 流程设计", 2)
    add_para(doc, "RAG 流程由文档入库和问题回答两部分组成。入库阶段将原始文档转换为文本块并建立向量索引；回答阶段将用户问题向量化，检索相似文本块，构建上下文提示词并调用大模型生成答案。")
    add_ascii_diagram(doc, "图3-2 RAG 工作流程图", [
        "文档上传 → 文本抽取 → 文本分块 → 向量生成 → Faiss 索引持久化",
        "用户提问 → 查询向量 → TopK 检索 → 重排 → Prompt 构造 → LLM 回答"
    ])

    add_heading(doc, "第四章 系统实现", 1)
    add_heading(doc, "4.1 用户认证实现", 2)
    add_para(doc, "认证模块由 AuthService 和认证路由共同实现。密码使用 bcrypt 加盐哈希保存；登录成功后生成 JWT，页面路由通过 cookie 中的 token 判断用户状态；修改类请求通过 CSRF Token 校验。系统限制密码长度和复杂度，并记录登录失败次数，达到阈值时写入锁定时间。")
    add_heading(doc, "4.2 知识库与文档处理实现", 2)
    add_para(doc, "文档处理由 FileProcessor 实现。PDF 通过 PyMuPDF 读取页面文本，若页面文本过少则转为图片并使用 PaddleOCR 识别；DOCX 读取段落和表格；TXT 使用 UTF-8 方式读取。TextSplitter 对文本按页面标记和段落进行切分，并处理超长段落。")
    add_heading(doc, "4.3 向量检索与问答实现", 2)
    add_para(doc, "EmbeddingModel 优先调用 Ollama 的 /api/embeddings 接口，默认嵌入模型为 bge-m3；连接失败时使用字符哈希向量作为降级方案。VectorStore 使用 Faiss IndexFlatIP 建立内积索引，向量归一化后可近似计算余弦相似度。QAModule.query_with_knowledge 负责把检索、重排、提示词构造和模型调用串联起来。")
    add_heading(doc, "4.4 学习记录实现", 2)
    add_para(doc, "学习记录由 LearningRecord 模型保存，包括问题、答案、来源、会话编号、分类、知识库名称和创建时间。接口 /api/records/sessions 按 session_id 生成最近会话列表，并使用内存缓存减少重复查询；/api/feedback 可记录用户对回答的有用性反馈。")

    add_heading(doc, "第五章 系统测试", 1)
    add_heading(doc, "5.1 测试方法", 2)
    add_para(doc, "系统测试采用功能测试、异常测试和性能观察相结合的方法。功能测试覆盖注册登录、知识库创建、文档上传、智能问答和历史记录；异常测试覆盖文件格式、模型不可用、Token 过期和 CSRF 缺失；性能观察关注文档入库耗时、向量检索速度和会话列表缓存效果。")
    add_table(doc, "表5-1 测试用例摘要", ["测试对象", "测试内容", "通过标准"], [
        ["用户认证", "合法注册登录、错误密码锁定、密码重置", "状态变化正确，无未处理异常"],
        ["知识库", "创建、上传、统计、删除来源", "向量数量和文件来源正确更新"],
        ["RAG 问答", "围绕上传文档提问", "答案使用上下文且返回 sources"],
        ["历史记录", "多轮问答后查询会话", "会话按时间倒序显示"],
        ["异常处理", "上传非法格式、关闭模型服务", "返回友好错误信息"]
    ], [3.0, 7.0, 6.0])
    add_heading(doc, "5.2 测试结论", 2)
    add_para(doc, "测试方案能够覆盖系统核心业务链路。系统具备较完整的异常提示和安全控制，Faiss 本地检索可以满足课程设计规模的数据查询需求。后续仍可补充自动化接口测试和更大规模知识库下的压测。")

    add_heading(doc, "第六章 总结", 1)
    add_para(doc, "本课程设计完成了智学伴智能学习系统的需求分析、系统设计、开发实现和测试设计。系统的创新点在于将学生个人资料库与 RAG 问答结合，既能调用大模型生成自然语言回答，又能通过检索来源降低回答脱离资料的问题。通过本次设计，进一步掌握了 Flask Web 系统分层开发、数据库建模、文档解析、向量检索和大模型集成等知识。系统仍有改进空间，例如可引入更稳定的重排模型、完善权限角色、增加回答来源可视化和自动化测试覆盖率。")
    add_heading(doc, "参考文献", 1)
    refs = [
        "[1] Flask Documentation. Pallets Projects.",
        "[2] SQLAlchemy Documentation. SQLAlchemy Project.",
        "[3] Johnson J, Douze M, Jegou H. Billion-scale similarity search with GPUs. IEEE Transactions on Big Data, 2019.",
        "[4] Lewis P, Perez E, Piktus A, et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS, 2020.",
        "[5] Ollama Documentation. Ollama Project.",
        "[6] PaddleOCR Documentation. PaddlePaddle Project."
    ]
    for ref in refs:
        add_para(doc, ref, first=False)
    doc.add_page_break()
    add_heading(doc, "评定意见页", 1)
    for line in ["指导教师评语：", "", "", "", "成绩：____________", "指导教师签名：____________", "日期：____________"]:
        add_para(doc, line, first=False)
    return doc


def main():
    OUT_DIR.mkdir(exist_ok=True)
    docs = {
        "实验2_智能系统设计实验_已完成.docx": experiment2(),
        "实验3_智能系统开发实验_已完成.docx": experiment3(),
        "实验4_智能系统测试实验_已完成.docx": experiment4(),
        "智学伴智能学习系统_课程设计报告_已完成.docx": course_report(),
    }
    for name, doc in docs.items():
        doc.save(OUT_DIR / name)
        print(OUT_DIR / name)


if __name__ == "__main__":
    main()
