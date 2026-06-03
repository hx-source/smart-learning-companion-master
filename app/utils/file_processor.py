"""上传文档解析工具。

知识库只处理纯文本，所以这里负责把 PDF/TXT/DOCX 转成字符串。
PDF 会优先提取内嵌文本；如果页面像扫描件或只提取到页码，则使用 OCR 兜底。
"""

import os
from docx import Document
import fitz  # PyMuPDF
from paddleocr import PaddleOCR


class FileProcessor:
    """文件内容提取器。"""

    def __init__(self):
        # PaddleOCR 初始化较重，因此在处理器实例创建时统一准备好。
        self.ocr = PaddleOCR(use_angle_cls=True, lang='ch', show_log=False)

    def process_file(self, file_path):
        """根据文件扩展名调用对应的处理方法。"""
        file_extension = os.path.splitext(file_path)[1].lower()

        if file_extension == '.pdf':
            return self.process_pdf(file_path)
        elif file_extension == '.txt':
            return self.process_txt(file_path)
        elif file_extension == '.docx':
            return self.process_docx(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {file_extension}")

    def process_pdf(self, file_path):
        """处理 PDF 文件，提取文本。

        对普通 PDF 使用 PyMuPDF 直接取文本；对扫描页或文本过短的页面使用 OCR。
        """
        text = []
        try:
            doc = fitz.open(file_path)
            num_pages = doc.page_count
            for page_num in range(num_pages):
                page = doc[page_num]
                page_text = page.get_text()
                
                # 检查提取的文本是否只是页码。
                # 有些扫描 PDF 的文本层只有页码，直接使用会导致知识库几乎无内容。
                is_only_page_number = False
                if page_text:
                    stripped_text = page_text.strip()
                    # 检查是否只是数字（可能是页码）。
                    if stripped_text.isdigit() and len(stripped_text) <= 3:
                        is_only_page_number = True
                
                if page_text and len(page_text.strip()) > 10 and not is_only_page_number:
                    text.append(f"=== 第 {page_num + 1} 页 ===")
                    try:
                        page_text = page_text.encode('utf-8', errors='replace').decode('utf-8')
                    except Exception:
                        pass
                    text.append(page_text)
                else:
                    text.append(f"=== 第 {page_num + 1} 页 (OCR识别) ===")
                    ocr_text = self._ocr_page(page)
                    if ocr_text:
                        text.append(ocr_text)
                    else:
                        # 如果 OCR 也失败，使用原始文本，至少保留可提取到的信息。
                        if page_text:
                            text.append(page_text)
            doc.close()
            result = '\n'.join(text)
            try:
                result = result.encode('utf-8', errors='replace').decode('utf-8')
            except Exception:
                pass
            return result
        except Exception as e:
            raise Exception(f"处理PDF文件时出错: {str(e)}")

    def _ocr_page(self, page):
        """使用 PaddleOCR 识别 PDF 页面。"""
        try:
            # 尝试不同的缩放比例，提高识别精度；取文本量最多的一次作为结果。
            scales = [3, 4]
            best_result = []
            
            for scale in scales:
                mat = fitz.Matrix(scale, scale)
                pix = page.get_pixmap(matrix=mat, alpha=False)
                img_bytes = pix.tobytes("png")
                
                # OCR 输入使用渲染后的 PNG 字节，避免依赖临时图片文件。
                result = self.ocr.ocr(img_bytes, cls=True)
                
                if result and result[0]:
                    ocr_lines = []
                    for line in result[0]:
                        if line and len(line) > 1:
                            text = line[1][0]
                            if text and len(text.strip()) > 0:
                                ocr_lines.append(text)
                    if len(' '.join(ocr_lines)) > len(' '.join(best_result)):
                        best_result = ocr_lines
            
            # 如果彩色图识别不到内容，尝试灰度渲染，部分扫描件会更清晰。
            if not best_result:
                # 尝试灰度处理
                mat = fitz.Matrix(3, 3)
                pix = page.get_pixmap(matrix=mat, alpha=False, colorspace=fitz.csGRAY)
                img_bytes = pix.tobytes("png")
                
                result = self.ocr.ocr(img_bytes, cls=True)
                
                if result and result[0]:
                    for line in result[0]:
                        if line and len(line) > 1:
                            text = line[1][0]
                            if text and len(text.strip()) > 0:
                                best_result.append(text)
            
            return '\n'.join(best_result)
        except Exception as e:
            print(f"OCR识别出错: {str(e)}")
            return ""

    def process_txt(self, file_path):
        """处理 TXT 文件，提取文本。"""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='replace') as file:
                return file.read()
        except Exception as e:
            raise Exception(f"处理TXT文件时出错: {str(e)}")

    def process_docx(self, file_path):
        """处理 DOCX 文件，提取段落和表格文本。"""
        text = []
        try:
            doc = Document(file_path)
            # 处理段落：Word 正文的大多数内容都在 paragraphs 中。
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text.append(paragraph.text)
            # 处理表格：把同一行的单元格用制表符连接，保留基本结构。
            for table in doc.tables:
                for row in table.rows:
                    row_text = '\t'.join([cell.text for cell in row.cells])
                    if row_text:
                        text.append(row_text)
            return '\n'.join(text)
        except Exception as e:
            raise Exception(f"处理DOCX文件时出错: {str(e)}")
